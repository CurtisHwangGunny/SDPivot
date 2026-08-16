import io
import unittest

from pptx import Presentation

from docreader.parser.markitdown_parser import MarkitdownParser
from docreader.parser.parser import Parser
from docreader.parser.registry import BUILTIN_ENGINE, registry


def _minimal_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Builtin PowerPoint"
    slide.placeholders[1].text = "PowerPoint content parsed as markdown"

    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


class TestBuiltinPptParser(unittest.TestCase):
    def test_builtin_registry_resolves_pptx_and_ppt(self):
        self.assertIs(
            registry.get_parser_class("", "pptx"),
            MarkitdownParser,
        )
        self.assertIs(
            registry.get_parser_class("", "ppt"),
            MarkitdownParser,
        )

    def test_builtin_engine_metadata_includes_pptx_and_ppt(self):
        builtin = next(
            engine
            for engine in registry.list_engines()
            if engine["name"] == BUILTIN_ENGINE
        )

        self.assertIn("pptx", builtin["file_types"])
        self.assertIn("ppt", builtin["file_types"])

    def test_builtin_engine_parses_pptx_to_non_empty_markdown(self):
        document = Parser().parse_file(
            file_name="minimal.pptx",
            file_type="pptx",
            content=_minimal_pptx(),
            parser_engine="",
        )

        self.assertTrue(document.content.strip())
        self.assertIn("Builtin PowerPoint", document.content)
        self.assertIn("PowerPoint content parsed as markdown", document.content)


if __name__ == "__main__":
    unittest.main()
